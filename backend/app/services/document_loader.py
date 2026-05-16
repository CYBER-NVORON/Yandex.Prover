from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from app.config import MAX_CHARS_FOR_ANALYSIS


class DocumentLoadError(Exception):
    """Base error for user-facing document loading failures."""


class UnsupportedFormatError(DocumentLoadError):
    pass


class EmptyDocumentError(DocumentLoadError):
    pass


@dataclass(frozen=True)
class LoadedDocument:
    filename: str
    text: str
    extension: str
    was_truncated: bool = False
    original_chars: int = 0


SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".pptx"}


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentLoadError("Не удалось распознать кодировку текстового файла.")


def _extract_pdf(content: bytes) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(content))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(page.strip() for page in pages if page.strip())
    except Exception as exc:  # pragma: no cover - exact parser errors vary
        raise DocumentLoadError("Не удалось извлечь текст из PDF.") from exc


def _extract_docx(content: bytes) -> str:
    try:
        from docx import Document

        doc = Document(BytesIO(content))
        paragraphs = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
        return "\n".join(paragraphs)
    except Exception as exc:  # pragma: no cover
        raise DocumentLoadError("Не удалось извлечь текст из DOCX.") from exc


def _extract_pptx(content: bytes) -> str:
    try:
        from pptx import Presentation

        presentation = Presentation(BytesIO(content))
        chunks: list[str] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            slide_text: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    slide_text.append(shape.text.strip())
            if slide_text:
                chunks.append(f"Слайд {slide_number}: " + "\n".join(slide_text))
        return "\n\n".join(chunks)
    except Exception as exc:  # pragma: no cover
        raise DocumentLoadError("Не удалось извлечь текст из PPTX.") from exc


def _truncate(text: str, max_chars: int | None) -> tuple[str, bool, int]:
    original_chars = len(text)
    if max_chars is None:
        return text, False, original_chars
    if original_chars <= max_chars:
        return text, False, original_chars
    return text[:max_chars], True, original_chars


def load_document_from_bytes(
    filename: str,
    content: bytes,
    *,
    max_chars: int | None = None,
) -> LoadedDocument:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise UnsupportedFormatError(f"Формат {extension or 'без расширения'} не поддерживается. Поддерживаются: {supported}.")
    if not content:
        raise EmptyDocumentError("Файл пустой.")

    if extension in {".txt", ".md"}:
        text = _decode_text(content)
    elif extension == ".pdf":
        text = _extract_pdf(content)
    elif extension == ".docx":
        text = _extract_docx(content)
    elif extension == ".pptx":
        text = _extract_pptx(content)
    else:  # pragma: no cover - guarded above
        raise UnsupportedFormatError("Формат не поддерживается.")

    text = text.strip()
    if not text:
        raise EmptyDocumentError("Не удалось найти текст в файле.")

    text, was_truncated, original_chars = _truncate(text, max_chars)
    return LoadedDocument(
        filename=filename,
        text=text,
        extension=extension,
        was_truncated=was_truncated,
        original_chars=original_chars,
    )


def load_document(path: str | Path, *, max_chars: int | None = None) -> LoadedDocument:
    file_path = Path(path)
    try:
        content = file_path.read_bytes()
    except OSError as exc:
        raise DocumentLoadError(f"Не удалось прочитать файл: {file_path}") from exc
    return load_document_from_bytes(file_path.name, content, max_chars=max_chars)


